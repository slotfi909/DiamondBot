import re
import logging
from openpyxl import Workbook
from io import BytesIO
from datetime import datetime
from typing import List, TypeAlias
from types import FunctionType
from telegram import Update
from telegram.ext import ContextTypes
from prisma.enums import UserRole, Grade, Field, PaymentStatus, PaymentMethod
from prisma import Prisma

# Constants
from src.constants import Matcher, PENDING_ORDER_LIFETIME_MINUTES, MERCHANT_ID

# Utils
from src.utils.message import send_message
from src.utils.db import db
from src.utils.user import get_user
from src.utils.time import convert_to_jajali, to_utc, substract_minutes
from src.payment.zarinpal import Zarinpal

Matchers: TypeAlias = List[Matcher]

zarinpal = Zarinpal(MERCHANT_ID)


def create_matcher(command: str, callback: FunctionType) -> Matcher:
    return {
        "command": command,
        "callback": callback
    }


# This function is a matcher when the bot get started by a link with
# special args that we use for joining chat and referral link
# it uses the first argument and the rest of arguments get passed to matchers
def start_arg_matcher(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    message_sender = send_message(update, ctx)

    async def runner(matchers: Matchers):
        regex = r'^[^-]+-[^-]+$'

        first_arg = ctx.args[0]

        if not re.match(regex, first_arg):
            return

        command, value = first_arg.split("-")

        ctx.user_data["matcher"] = {
            "command": command,
            "value": value
        }

        for matcher in matchers:
            if matcher["command"] == command:
                return await matcher["callback"](update, ctx, message_sender)

    return runner


def compact_price(price: int):
    stringified_price = str(price)[::-1]

    final_price = ''

    for i in range(1, len(stringified_price) + 1):
        final_price += stringified_price[i - 1]

        if i != 0 and i % 3 == 0 and i != len(stringified_price):
            final_price += ','

    return final_price[::-1]


async def get_refund_count(tel_id: int, tx=None):
    query_result = await (tx or db).query_raw(f"""--sql
        SELECT COUNT(*)
        FROM "Order" AS o
            INNER JOIN "User" as u ON (u.user_id = o.user_id)
            INNER JOIN "Payment" as p ON (p.payment_id = o.payment_id)
        WHERE u.tel_id = $1 AND p.status = 'COMPLETED' AND o.refunded_at IS NOT NULL
      """, tel_id)

    return query_result[0]["count"]


async def get_advisor_id(tel_id: int):
    advisors_count = await db.user.count(where={
        "OR": [
            {
                "user_role": UserRole.ADVISOR
            },
            {
                "user_role": UserRole.ADMIN
            },
        ]
    })

    return (tel_id % advisors_count) + 10


def get_real_grade(grade: Grade):
    real_grade = ''

    match grade:
        case Grade.TENTH:
            real_grade = "دهم"
        case Grade.ELEVENTH:
            real_grade = "یازدهم"
        case Grade.TWELVTH:
            real_grade = "دوازدهم"
        case Grade.DIPLOMA:
            real_grade = "پشت کنکوری"

    return real_grade


def get_real_field(field: Field):
    real_field = ''

    match field:
        case Field.MATH:
            real_field = "ریاضی"
        case Field.MEDICAL:
            real_field = "تجربی"
        case Field.HUMANITIES:
            real_field = "انسانی"

    return real_field


async def get_user_info(tel_id: int):
    user = await get_user(tel_id)

    jalali_date = convert_to_jajali(user.dob)

    return (
        f"اطلاعات دانش آموز با شناسه {user.user_id}:\n\n"
        f"نام: {user.firstname} {user.lastname}\n"
        f"کد ملی: {user.national_code}\n"
        f"شماره تماس: {user.phone_number}\n"
        f"تاریخ تولد: {
            jalali_date.year}/{jalali_date.month}/{jalali_date.day}\n"
        f"پایه تحصیلی: {get_real_grade(user.grade)}\n"
        f"رشته تحصیلی: {get_real_field(user.field)}\n"
    )


async def build_customer_details_excel(update: Update, ctx: ContextTypes.DEFAULT_TYPE, product_id: int):
    customers = await db.query_raw(f"""--sql
        SELECT u.firstname, u.lastname, u.national_code, u.phone_number, u.dob, u.grade, u.field, p.title, pay.amount,
               u.tel_chat_id
        FROM "User" AS u
            INNER JOIN "Order" AS o ON (o.user_id = u.user_id)
            INNER JOIN "Product" AS p ON (o.product_id = p.product_id)
            INNER JOIN "Payment" AS pay ON (o.payment_id = pay.payment_id)
        WHERE p.product_id = $1 AND pay.status = '{str(PaymentStatus.COMPLETED).upper()}'
        ORDER BY u.created_at desc
    """, product_id)

    wb = Workbook()
    ws = wb.active

    headers = ["اسم کوچک", "نام خانوادگی", "کد ملی", "شماره تماس", "تاریخ تولد",
               "پایه تحصیلی", "رشته تحصیلی", "بسته خریداری شده", "مبلغ پرداخت شده", "username تلگرام"]

    for i, header in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=header)

    for i, customer in enumerate(customers, 1):
        customer_chat = None

        try:
            customer_chat = await ctx.bot.get_chat(chat_id=customer["tel_chat_id"])
        except Exception as ex:
            logging.error(ex)

        customer_username = customer_chat.username if customer_chat else None

        customer.pop("tel_chat_id")

        customer["username"] = f"@{
            customer_username}" if customer_username else "نامشخص"

        for j, (key, value) in enumerate(customer.items(), 1):
            real_value = value

            if key == 'dob':
                jalali_date = convert_to_jajali(datetime.fromisoformat(value))

                real_value = f"{
                    jalali_date.year}/{jalali_date.month}/{jalali_date.day}"

            if key == 'grade':
                real_value = get_real_grade(value)

            if key == 'field':
                real_value = get_real_field(value)

            ws.cell(row=i + 1, column=j, value=real_value)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer


async def create_order(product_id: int, target_user_tel_id: int, payment_method: PaymentMethod, payment_status: PaymentStatus = PaymentStatus.PENDING, amount: int | None = None, refunded_at: datetime | None = None, outter_tx: Prisma | None = None):
    async with db.tx() as tx:
        product = await (outter_tx or tx).product.find_unique(where={
            "product_id": product_id
        })

        query_result = await (outter_tx or tx).query_raw(f"""--sql
            SELECT COUNT(*)
            FROM "Order" AS o
                INNER JOIN "Payment" AS p ON (p.payment_id = o.payment_id)
            WHERE o.product_id = $1 AND (
                (p.status = '{str(PaymentStatus.PENDING).upper()}' AND o.created_at >= $2::timestamp with time zone AT TIME zone 'UTC')
                OR
                p.status = '{str(PaymentStatus.COMPLETED).upper()}'
            )
        """, product_id, to_utc(substract_minutes(to_utc(datetime.now()), PENDING_ORDER_LIFETIME_MINUTES)))

        bought_count = query_result[0]["count"]

        if product.capacity is not None and product.capacity <= bought_count:
            return product, None, False

        order = await (outter_tx or tx).order.create(data={
            "product": {
                "connect": {
                    "product_id": product_id
                }
            },
            "user": {
                "connect": {
                    "tel_id": target_user_tel_id
                }
            },
            "refunded_at": refunded_at,
            "payment": {
                "create": {
                    "amount": amount or product.price,
                    "method": payment_method,
                    "status": payment_status,
                }
            }
        }, include={
            "payment": True
        })

    return product, order, True


def to_rial(toman: int):
    return toman * 10
